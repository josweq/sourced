"""Convierte las hojas de revisión humana en las etiquetas que usa la capa de IA.

Uso:
    python scripts/consolidar_etiquetas.py data/local/etiquetado/Etiquetado-Lupa-*.xlsx
Escribe evaluation/etiquetas/temas.csv y pares.csv. Reglas:
- «Sí» en «¿De acuerdo?» adopta la propuesta de Codex; «No» usa «Tu tema» (temas) o invierte la propuesta (pares).
- Una fila sin revisar NO se acepta: queda fuera (las etiquetas deben ser humanas).
- Se registra quién revisó cada fila (nombre de la pestaña).
"""
import csv
from pathlib import Path
import sys

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SALIDA = ROOT / "evaluation" / "etiquetas"
TEMAS = {"economia", "logistica_canal", "turismo", "servicios_publicos", "eventos_naturales", "regulacion", "sin_clasificar"}


def leer(path):
    wb = load_workbook(path, data_only=True)
    hoja = next(ws for ws in wb.worksheets if ws.title != "Guía")
    persona = hoja.title
    temas, pares, seccion = [], [], None
    for fila in hoja.iter_rows(values_only=True):
        celdas = ["" if v is None else str(v).strip() for v in fila]
        if celdas[0] == "id_noticia":
            seccion = "temas"; continue
        if celdas[0] == "id_a":
            seccion = "pares"; continue
        if not celdas[0].startswith("N-"):
            continue
        if seccion == "temas":
            _id, titulo, _medio, propuesta, _conf, _motivo, acuerdo, tu_tema = celdas[:8]
            if acuerdo == "Sí" and propuesta in TEMAS:
                temas.append({"id_noticia": _id, "titulo": titulo, "etiqueta_humana": propuesta, "revisor": persona, "fuente": "propuesta aceptada"})
            elif acuerdo == "No" and tu_tema in TEMAS:
                temas.append({"id_noticia": _id, "titulo": titulo, "etiqueta_humana": tu_tema, "revisor": persona, "fuente": "corregida"})
        elif seccion == "pares":
            id_a, _ta, id_b, _tb, propuesta, _conf, acuerdo = celdas[:7]
            if acuerdo in {"Sí", "No"} and propuesta in {"Sí", "No"}:
                mismo = propuesta if acuerdo == "Sí" else ("No" if propuesta == "Sí" else "Sí")
                pares.append({"id_a": id_a, "id_b": id_b, "mismo_evento": 1 if mismo == "Sí" else 0, "revisor": persona})
    return persona, temas, pares


def main():
    archivos = [Path(p) for p in sys.argv[1:]]
    if not archivos:
        raise SystemExit("Indica las hojas revisadas (.xlsx)")
    todos_t, todos_p = [], []
    for path in archivos:
        persona, temas, pares = leer(path)
        print(f"{persona}: {len(temas)} temas revisados, {len(pares)} pares revisados")
        todos_t += temas; todos_p += pares
    SALIDA.mkdir(parents=True, exist_ok=True)
    with (SALIDA / "temas.csv").open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, ["id_noticia", "titulo", "etiqueta_humana", "revisor", "fuente"]); w.writeheader(); w.writerows(todos_t)
    with (SALIDA / "pares.csv").open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, ["id_a", "id_b", "mismo_evento", "revisor"]); w.writeheader(); w.writerows(todos_p)
    corregidas = sum(1 for t in todos_t if t["fuente"] == "corregida")
    print(f"Total: {len(todos_t)} temas ({corregidas} corregidos por persona), {len(todos_p)} pares → {SALIDA}")


if __name__ == "__main__":
    main()
